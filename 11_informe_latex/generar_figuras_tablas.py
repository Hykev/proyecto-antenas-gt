"""Genera figuras y tablas .tex del informe a partir de CSV y del Excel (valores recalculados)."""
import os, sys, json, math, csv
import numpy as np, pandas as pd, matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE); os.chdir(HERE)
os.makedirs('tablas', exist_ok=True); FIG = 'figuras/'
plt.rcParams.update({'font.size': 9})
UNI = {'^': '\\textasciicircum{}', '·': '$\\cdot$', '<': '$<$', '>': '$>$', '≈': '$\\approx$', '≤': '$\\leq$', '≥': '$\\geq$', '−': '-', '↔': '$\\leftrightarrow$', '→': '$\\rightarrow$', '↓': 'DL', '↑': 'UL', '²': '$^2$', '₂': '$_2$', 'Ω': 'Ohm', '“': '"', '”': '"', '’': "'", 'µ': '$\\mu$'}
def tx(s):
    s = str(s).replace('\\', '').replace('%', '\\%').replace('&', '\\&').replace('_', '\\_').replace('#', '\\#').replace('$', '\\$')
    for a, b in UNI.items(): s = s.replace(a, b)
    return s
df = pd.read_csv('../04_propagacion_linkbudget/linkbudget_por_casa.csv')
# ---------- diagrama del sistema ----------
def box(ax, x, y, w, h, t, c='#dbe9f7', fs=8):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0.02', fc=c, ec='#1F4E78', lw=1)); ax.text(x + w / 2, y + h / 2, t, ha='center', va='center', fontsize=fs)
def arrow(ax, a, b, t=''):
    ax.annotate('', xy=b, xytext=a, arrowprops=dict(arrowstyle='->', color='#333', lw=1.2));
    if t: ax.text((a[0] + b[0]) / 2, (a[1] + b[1]) / 2 + .12, t, ha='center', fontsize=7)
fig, ax = plt.subplots(figsize=(10, 4.6)); ax.set_xlim(0, 10); ax.set_ylim(0, 4.6); ax.axis('off')
box(ax, .1, 3.2, 1.6, .9, 'Constelación\nStarlink', '#e8e8e8'); box(ax, 2.2, 3.2, 1.9, .9, 'Plato Starlink\nPerformance Gen 3\n(torre, 30 m)', '#fff2cc')
box(ax, 2.2, 1.8, 1.9, .9, 'Router + switch PoE\n(gabinete, torre)', '#dbe9f7'); box(ax, 2.2, .3, 1.9, .9, 'Solar: 3×550 W\n+ 2 baterías 5.12 kWh', '#e2f0d9')
for i, (s, ch) in enumerate((('S1', '5745'), ('S2', '5785'), ('S3', '5825'))): box(ax, 5.0, 3.3 - i * 1.15, 2.0, .85, f'Rocket Prism + AM-5G16-120\nSector {s} · {ch} MHz', '#dbe9f7', 7)
box(ax, 7.8, 3.3, 2.0, .85, 'CPE LiteBeam 5AC\n(38 casas, S1)', '#fce4d6'); box(ax, 7.8, 2.15, 2.0, .85, 'CPE LiteBeam 5AC\n(39 casas, S2)', '#fce4d6'); box(ax, 7.8, 1.0, 2.0, .85, 'CPE LiteBeam 5AC\n(32 casas, S3)', '#fce4d6')
arrow(ax, (1.7, 3.65), (2.2, 3.65), 'Ku/Ka'); arrow(ax, (3.15, 3.2), (3.15, 2.7), 'Ethernet'); arrow(ax, (4.1, 2.25), (5.0, 3.7)); arrow(ax, (4.1, 2.25), (5.0, 2.55)); arrow(ax, (4.1, 2.25), (5.0, 1.4)); arrow(ax, (3.15, 1.2), (3.15, 1.8), 'CC 48 V')
for i in range(3): arrow(ax, (7.0, 3.7 - i * 1.15), (7.8, 3.7 - i * 1.15 - (0 if i == 0 else 0.0) - (0 if i == 0 else -0.0)))
ax.text(5, .25, 'Cada CPE: router Wi-Fi en la casa · 2 puntos de acceso comunitarios', fontsize=8, style='italic')
plt.tight_layout(); plt.savefig(FIG + 'fig_diagrama_sistema.png', dpi=200); plt.close()
# ---------- arquitectura B ----------
rng = np.random.default_rng(3); fig, ax = plt.subplots(figsize=(6.5, 3.4))
pts = np.array([(-350, 20), (-250, 90), (-250, -80), (-120, 40), (-120, -70), (0, 0), (120, 50), (120, -60), (250, 40), (350, -20)])
for i in range(len(pts)):
    for j in range(i + 1, len(pts)):
        if np.hypot(*(pts[i] - pts[j])) < 170: ax.plot(*zip(pts[i], pts[j]), color='#999', lw=.8)
ax.scatter(*pts.T, s=70, c='#3182bd', zorder=3); ax.scatter(0, 0, s=170, marker='^', c='r', zorder=4, label='Nodo con Starlink (pasarela)')
ax.set(xlabel='Este (m)', ylabel='Norte (m)', title='Alternativa B: malla Wi-Fi exterior (10 nodos en postes)', aspect='equal'); ax.legend(fontsize=7, loc='lower left'); plt.tight_layout(); plt.savefig(FIG + 'fig_arquitectura_B.png', dpi=200); plt.close()
# ---------- perfiles ----------
prof = json.load(open('../04_propagacion_linkbudget/perfiles_google_earth.json')); fig, ax = plt.subplots(figsize=(8, 3.8))
for p in prof:
    d, z = zip(*p['puntos_dist_desde_torre_m_elev_m']); ax.plot(d, z, marker='o', ms=3, label=f"{p['perfil']}: casa {p['casa']} (az {p['azimut_torre_a_casa_deg']:.0f}°)")
ax.plot([0, 0], [206, 236], 'r', lw=3); ax.text(5, 232, 'Torre 30 m', color='r', fontsize=8)
ax.set(xlabel='Distancia desde la torre (m)', ylabel='Elevación (m s. n. m.)', title='Perfiles de terreno torre–casa (Google Earth)'); ax.legend(fontsize=7); ax.grid(alpha=.3); plt.tight_layout(); plt.savefig(FIG + 'fig_perfiles_terreno.png', dpi=200); plt.close()
# ---------- follaje vs altura ----------
H = [12, 18, 24, 30, 36]; n30 = [19, 18, 11, 4, 2]; med = [17.1, 17.1, 15.2, 12.0, 10.8]; p90 = [33.6, 33.6, 29.9, 27.4, 24.8]; mx = [54.3, 54.3, 45.7, 41.2, 37.8]
fig, ax = plt.subplots(1, 2, figsize=(9, 3.2)); ax[0].bar([str(h) for h in H], n30, color='#e6550d'); ax[0].set(xlabel='Altura de torre (m)', ylabel='Casas con > 30 dB de follaje', title='Casas críticas vs altura')
for k, v in (('Mediana', med), ('p90', p90), ('Máxima', mx)): ax[1].plot(H, v, marker='o', label=k)
ax[1].set(xlabel='Altura de torre (m)', ylabel='Pérdida por follaje (dB)', title='Pérdida Weissberger a 5.8 GHz'); ax[1].legend(); ax[1].grid(alpha=.3); plt.tight_layout(); plt.savefig(FIG + 'fig_follaje_vs_altura.png', dpi=200); plt.close()
# ---------- tablas desde datos ----------
# 1) sectores
rows = []
cap = {'S1': (.437, .177), 'S2': (.484, .206), 'S3': (.340, .131)}
for s in ('S1', 'S2', 'S3'):
    d = df[df.sector == s]; rows.append((s, len(d), f'{d.dist_m.max():.0f}', f'{d.L_follaje_dB.median():.1f}', f'{d.margen_dl_BPSK_dB.min():.1f}', f'{d.margen_ul_BPSK_dB.min():.1f}', f'{100*cap[s][0]:.0f}\\,\\% / {100*cap[s][1]:.0f}\\,\\%'))
open('tablas/t_sectores.tex', 'w').write('\\begin{tabular}{lrrrrrr}\\toprule Sector & Casas & Dist. máx. (m) & Follaje mediano (dB) & Margen mín. DL (dB) & Margen mín. UL (dB) & Util. aire DL/UL\\\\\\midrule\n' + '\\\\\n'.join(' & '.join(map(str, r)) for r in rows) + '\\\\\\bottomrule\\end{tabular}\n')
# 2) rangos
bins = [(0, 100), (100, 200), (200, 300), (300, 400), (400, 500)]; rows = []
for a, b in bins:
    d = df[(df.dist_m >= a) & (df.dist_m < b)]
    if len(d): rows.append((f'{a}--{b}', len(d), f'{d.G_ap_ef_dBi.median():.1f}', f'{d.FSPL_dB.median():.1f}', f'{d.L_follaje_dB.median():.1f}', f'{d.margen_dl_BPSK_dB.median():.1f}', f'{d.margen_ul_BPSK_dB.min():.1f}', tx(d.MCS_dl.mode()[0])))
open('tablas/t_rangos.tex', 'w').write('\\begin{tabular}{lrrrrrrl}\\toprule Rango (m) & Casas & $G_{AP}$ efec. (dBi) & FSPL (dB) & Follaje (dB) & Margen DL med. (dB) & Margen UL mín. (dB) & MCS DL más común\\\\\\midrule\n' + '\\\\\n'.join(' & '.join(map(str, r)) for r in rows) + '\\\\\\bottomrule\\end{tabular}\n')
# 3) peores casos y casos tipo
sel = []
for s in ('S1', 'S2', 'S3'):
    d = df[df.sector == s]; sel.append((f'Peor {s}', d.loc[d.margen_ul_BPSK_dB.idxmin()]))
sel.append(('Casa más cercana', df.loc[df.dist_m.idxmin()])); sel.append(('Casa mediana', df.iloc[(df.dist_m - df.dist_m.median()).abs().argsort().iloc[0]]))
lines = ['\\begin{tabular}{lrrrrrrrrrr}\\toprule Caso & Casa & Dist. (m) & $\\Delta$az (°) & $\\Delta$el (°) & $G_{AP}$ ef. & FSPL & Follaje & Margen DL & Margen UL & MCS DL \\\\\\midrule']
for n, r in sel: lines.append(f'{n} & {int(r.casa)} & {r.dist_m:.0f} & {r.dAz_deg:.1f} & {r.dEl_deg:.1f} & {r.G_ap_ef_dBi:.1f} & {r.FSPL_dB:.1f} & {r.L_follaje_dB:.1f} & {r.margen_dl_BPSK_dB:.1f} & {r.margen_ul_BPSK_dB:.1f} & {tx(r.MCS_dl)}\\\\')
open('tablas/t_casos.tex', 'w').write('\n'.join(lines) + '\n\\bottomrule\\end{tabular}\n')
# 4) MCS usados
vc = df.MCS_dl.value_counts(); order = ['256QAM 5/6', '256QAM 3/4', '64QAM 5/6', '64QAM 3/4', '64QAM 2/3', '16QAM 3/4', '16QAM 1/2']
open('tablas/t_mcs.tex', 'w').write('\\begin{tabular}{lr}\\toprule MCS (bajada, DL) & Casas\\\\\\midrule\n' + '\\\\\n'.join(f'{tx(m)} & {int(vc.get(m, 0))}' for m in order if vc.get(m, 0)) + '\\\\\\bottomrule\\end{tabular}\n')
# 5) BOM desde Excel
wb = load_workbook('../12_excel_memoria/Memoria_Calculo_Babilonia.xlsx', data_only=True); w = wb['BOM_CAPEX']; L = ['\\begin{longtable}{p{5.2cm}rrrrr}\\toprule Ítem & Cant. & P.\\,unit. (Q) & Total (Q) & Estado\\\\\\midrule\\endhead']
cats = {}
for r in range(5, 27):
    if w.cell(r, 1).value is None: continue
    n, cat, q, pu, tot, st = w.cell(r, 1).value, w.cell(r, 2).value, w.cell(r, 3).value, w.cell(r, 7).value, w.cell(r, 8).value, w.cell(r, 10).value
    cats[cat] = cats.get(cat, 0) + tot; est = 'Fuente' if str(st).startswith('Fuente') else 'Estim.'
    L.append(f'{tx(n)} & {int(q)} & {pu:,.0f} & {tot:,.0f} & {est}\\\\'.replace(',', '\\,'))
for r in range(27, 40):
    if w.cell(r, 1).value and str(w.cell(r, 1).value).startswith(('Subtotal', 'Flete', 'Imprev', 'CAPEX TOTAL (Q)', 'CAPEX por hogar')) and w.cell(r, 8).value is not None:
        L.append(f'\\midrule \\textbf{{{tx(w.cell(r, 1).value)}}} & & & \\textbf{{{w.cell(r, 8).value:,.0f}}} & \\\\'.replace(',', '\\,'))
open('tablas/t_bom.tex', 'w').write('\n'.join(L) + '\n\\bottomrule\\end{longtable}\n')
tot = sum(cats.values()); capx = wb['Resumen']['B19'].value
fig, ax = plt.subplots(1, 2, figsize=(9, 3.4)); ks = sorted(cats, key=cats.get, reverse=True); ax[0].barh(ks, [cats[k] / 1000 for k in ks], color='#3182bd'); ax[0].invert_yaxis(); ax[0].set(xlabel='Miles de Q (antes de flete e imprevistos)', title='CAPEX por categoría')
o = wb['OPEX_Tarifa']; lab = [o.cell(r, 1).value for r in range(11, 18)]; val = [o.cell(r, 3).value for r in range(11, 18)]
ax[1].barh(range(7), val, color='#e6550d'); ax[1].set_yticks(range(7)); ax[1].set_yticklabels([l[:28] for l in lab], fontsize=7); ax[1].invert_yaxis(); ax[1].set(xlabel='Q por mes (año 3)', title='OPEX mensual'); plt.tight_layout(); plt.savefig(FIG + 'fig_costos.png', dpi=200); plt.close()
# 6) OPEX tabla
L = ['\\begin{tabular}{lrr}\\toprule Concepto (Q/mes) & Año 1 & Año 3\\\\\\midrule']
for r in range(5, 27):
    n = o.cell(r, 1).value
    if r in (11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22): 
        a, b = o.cell(r, 2).value, o.cell(r, 3).value; f = (lambda v: f'{v:,.1f}' if r in (20, 21, 22) else f'{v:,.0f}')
        L.append(('\\midrule ' if r in (18, 22) else '') + f'{tx(n)} & {f(a)} & {f(b)}\\\\'.replace(',', '\\,'))
open('tablas/t_opex.tex', 'w').write('\n'.join(L) + '\n\\bottomrule\\end{tabular}\n')
# 7) sensibilidad
sv = pd.read_csv('../04_propagacion_linkbudget/sensibilidad_escenarios.csv'); L = ['\\begin{tabular}{p{6.4cm}rrrr}\\toprule Escenario & Casas útiles & Margen DL mín. & Margen UL mín. & Casas $>$30 dB\\\\\\midrule']
for _, r in sv.iterrows(): L.append(f'{tx(r.iloc[0])} & {int(r.iloc[1])} & {r.iloc[2]:.1f} & {r.iloc[3]:.1f} & {int(r.iloc[6])}\\\\')
open('tablas/t_sens.tex', 'w').write('\n'.join(L) + '\n\\bottomrule\\end{tabular}\n')
# 8) fuentes
fu = [l.strip().strip('|').split('|') for l in open('../14_fuentes/FUENTES.md', encoding='utf-8') if l.startswith('| F')]
L = ['\\begin{longtable}{p{.9cm}p{5.6cm}p{6.2cm}}\\toprule ID & Dato & Fuente\\\\\\midrule\\endhead']
import re
for r in fu:
    src = r[2].strip(); m = re.search(r'https?://[^\s,;)]+', src); link = f'\\url{{{m.group(0).rstrip(".")}}}' if m else tx(src[:70])
    L.append(f'{r[0].strip()} & {tx(r[1].strip()[:140])} & {link}\\\\')
open('tablas/t_fuentes.tex', 'w').write('\n'.join(L) + '\n\\bottomrule\\end{longtable}\n')
# anexo: tabla completa de casas
L = ['\\begin{longtable}{rlrrrrrrl}\\toprule Casa & Sec. & Dist. (m) & Elev. rayo (°) & Follaje (m) & Follaje (dB) & Margen DL & Margen UL & MCS DL\\\\\\midrule\\endhead']
for _, r in df.iterrows(): L.append(f'{int(r.casa)} & {r.sector} & {r.dist_m:.0f} & {r.el_deg:.1f} & {int(r.follaje_m)} & {r.L_follaje_dB:.1f} & {r.margen_dl_BPSK_dB:.1f} & {r.margen_ul_BPSK_dB:.1f} & {tx(r.MCS_dl)}\\\\')
open('tablas/t_anexo_casas.tex', 'w').write('\n'.join(L) + '\n\\bottomrule\\end{longtable}\n')
print('listo', capx, round(tot))
